# -*- coding: utf-8 -*-
"""Build DOCX with corrected homepage texts + short letter to contractor."""
from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_LINE_SPACING

PROJ = Path(__file__).resolve().parent
MD = PROJ / "Тексты_главной_Аллсан_ПРАВКИ_2026-09-11.md"
OUT = PROJ / "Тексты_главной_Аллсан_ПРАВКИ_2026-09-11.docx"
LETTER = PROJ / "Письмо_подрядчику_тексты_главной_2026-09-11.docx"


def set_run(run, *, bold=False, size=11, color=None):
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = color


def add_p(doc, text, *, bold=False, size=11, after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    r = p.add_run(text)
    set_run(r, bold=bold, size=size)
    return p


def md_to_docx(md_path: Path, out_path: Path) -> None:
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    for line in md_path.read_text(encoding="utf-8").splitlines():
        s = line.rstrip()
        if not s:
            add_p(doc, "", after=4)
            continue
        if s.startswith("# "):
            add_p(doc, s[2:], bold=True, size=16, after=10)
        elif s.startswith("## "):
            add_p(doc, s[3:], bold=True, size=13, after=8)
        elif s.startswith("### "):
            add_p(doc, s[4:], bold=True, size=12, after=6)
        elif s.startswith("**") and s.endswith("**") and s.count("**") == 2:
            add_p(doc, s.strip("*"), bold=True, size=11, after=4)
        elif s.startswith("- "):
            add_p(doc, "• " + s[2:], size=11, after=3)
        elif s.startswith("*(") and s.endswith(")*"):
            add_p(doc, s.strip("*()"), size=10, after=4)
        else:
            # light bold markers **x**
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(4)
            parts = s.split("**")
            for i, part in enumerate(parts):
                if not part:
                    continue
                r = p.add_run(part)
                set_run(r, bold=(i % 2 == 1), size=11)
    doc.save(out_path)


def build_letter(path: Path) -> None:
    doc = Document()
    add_p(doc, "Письмо подрядчику по текстам главной", bold=True, size=14, after=8)
    add_p(doc, "Дата: 11.09.2026", size=10, after=8)
    add_p(
        doc,
        "Добрый день! Спасибо за тексты главной — в целом они соответствуют согласованному смыслу "
        "и структуре. Ниже — правки, которые нужно внести перед фиксацией в макете.",
        after=8,
    )
    add_p(doc, "Принять как есть (без замечаний по смыслу)", bold=True, after=4)
    for x in [
        "Порядок блоков и hero (юнит-экономика / B2B / поддержка от 15 мин).",
        "Готовые интеграции → услуги → лицензии; Telegram-бот не включаем.",
        "ТОП-10 ЦРА, 600+ клиентов, с 2015 года.",
        "Кейсы Спектр / B2B / Электропром / Данахер; отзывы и видеоряд.",
        "СМИ «Промышленные страницы»; форма и контакты (комната 503).",
    ]:
        add_p(doc, "• " + x, after=2)
    add_p(doc, "Просим поправить", bold=True, after=4)
    for x in [
        "В русских текстах писать «ИТ», не «IT» (телеком и ИТ).",
        "Сертификат — «ИСО» (или с номером, если передадим).",
        "В преимуществах явно: отзывы, подтверждённые 1С.",
        "В витрине лицензий учесть категории из брифа: в т.ч. «Другие программы 1С» и "
        "«Дополнительные (пользовательские) лицензии»; CTA в каталог — «Смотреть все лицензии…», "
        "без отдельной категории каталога «Все лицензии».",
        "Кейс B2B — опора на https://alsn.ru/caseecomsc ; Данахер — лого можно.",
        "В FAQ добавить вопросы «Кто такая Аллсан Интеграция?» и «Вы официальный партнёр 1С / ЦРА?»; "
        "базу также брать с https://alsn.ru/about_us .",
        "В блоке «О компании» и ключевых заголовках использовать полное «Аллсан Интеграция» "
        "(короткое «Аллсан» допустимо в подписях).",
        "В письме/SEO-списке не использовать формулировку «1С Битрикс24» — правильно: "
        "«Битрикс24» / «внедрение Битрикс24 и интеграция с 1С».",
        "Отдельные SEO-посадочные Wildberries и Ozon — в структуре сайта; на главной оставляем "
        "общую карточку «Ozon и Wildberries», как сейчас.",
    ]:
        add_p(doc, "• " + x, after=2)
    add_p(
        doc,
        "Полная вычищенная версия текстов приложена: Тексты_главной_Аллсан_ПРАВКИ_2026-09-11.docx. "
        "Просьба перенести правки в рабочий документ и подтвердить.",
        after=8,
    )
    add_p(doc, "С уважением,", after=2)
    add_p(doc, "команда Аллсан Интеграции", after=2)
    doc.save(path)


if __name__ == "__main__":
    md_to_docx(MD, OUT)
    build_letter(LETTER)
    print("wrote", OUT)
    print("wrote", LETTER)
